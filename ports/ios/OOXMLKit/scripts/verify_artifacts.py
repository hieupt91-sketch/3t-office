#!/usr/bin/env python3
"""Kiểm chéo tệp do OOXMLKit lưu ra bằng bộ đọc ĐỘC LẬP (không dùng mã Swift):

1. zipfile.testzip(): giải nén mọi mục bằng zlib và kiểm CRC-32 (chứng minh bộ DEFLATE Swift đúng chuẩn).
2. Mọi phần .xml/.rels phân tích được bằng expat (xml.parsers.expat).
3. Mở bằng thư viện phổ biến: openpyxl (.xlsx), python-docx (.docx), python-pptx (.pptx) nếu đã cài.

Dùng:
  OOXMLKIT_ARTIFACTS=/tmp/out swift test
  python3 scripts/verify_artifacts.py /tmp/out [--json report.json]
Mã thoát 0 khi mọi tệp PASS.
"""
import json, os, sys, zipfile
import xml.parsers.expat


def check_xml(data):
    p = xml.parsers.expat.ParserCreate(namespace_separator=" ")
    p.Parse(data, True)


def open_with_library(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path)
            return f"openpyxl: {len(wb.sheetnames)} sheet, {sum(ws.max_row for ws in wb.worksheets)} hàng"
        if ext == ".docx":
            import docx
            d = docx.Document(path)
            return f"python-docx: {len(d.paragraphs)} đoạn, {len(d.tables)} bảng"
        if ext == ".pptx":
            import pptx
            p = pptx.Presentation(path)
            return f"python-pptx: {len(p.slides)} slide"
    except ImportError:
        return "NOT RUN (thiếu thư viện)"
    return "-"


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    root = sys.argv[1]
    results = []
    for name in sorted(os.listdir(root)):
        path = os.path.join(root, name)
        if not os.path.isfile(path):
            continue
        r = {"file": name, "status": "PASS", "detail": []}
        try:
            with zipfile.ZipFile(path) as z:
                bad = z.testzip()
                if bad:
                    raise RuntimeError(f"CRC/giải nén lỗi ở {bad}")
                n = 0
                for info in z.infolist():
                    if info.filename.endswith((".xml", ".rels")):
                        check_xml(z.read(info))
                        n += 1
                r["detail"].append(f"zlib+CRC OK {len(z.infolist())} mục, expat OK {n} XML")
            if name.lower().endswith((".xlsx", ".docx", ".pptx")):
                r["detail"].append(open_with_library(path))
        except Exception as e:  # noqa: BLE001 — báo cáo mọi lỗi, không nuốt
            r["status"] = "FAIL"
            r["detail"].append(f"{type(e).__name__}: {e}")
        results.append(r)
        print(f'{r["status"]:4}  {name:48}  ' + " | ".join(r["detail"]))
    failed = [r for r in results if r["status"] != "PASS"]
    print(f"\nTổng: {len(results)} tệp, PASS {len(results) - len(failed)}, FAIL {len(failed)}")
    if "--json" in sys.argv:
        with open(sys.argv[sys.argv.index("--json") + 1], "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=1)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
