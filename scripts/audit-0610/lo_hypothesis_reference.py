"""Giá trị tham chiếu thứ hai (LibreOffice) cho T.TEST/F.TEST/CHISQ.TEST/GAMMA khi Excel COM không dùng được.
Ghi XLSX có công thức (tiền tố _xlfn. như Excel), để soffice tính rồi xuất CSV. In ra: công thức -> giá trị.
Dùng làm bộ giá trị cho tests/unit/formula-hypothesis-tests.test.js."""
import csv, os, pathlib, re, subprocess, sys

import openpyxl

OUT = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else r"E:\VNOffice-profiles\claude-lo-ref")
OUT.mkdir(parents=True, exist_ok=True)
so = os.path.expanduser(r"~\tools\libreoffice\program\soffice.exe")
data = {
    "A": [3, 5, 7, 9, 11, 14], "B": [2, 6, 8, 8, 12, 13], "C": [4, 7, 9],
    "E": [10, 30], "F": [20, 40], "G": [12, 28], "H": [18, 42],
    "I": [5, 7, 9, 11], "J": [7, 7, 9, 9],
}
formulas = [
    "T.TEST(A1:A6,B1:B6,2,1)", "T.TEST(A1:A6,B1:B6,1,1)", "T.TEST(A1:A6,C1:C3,2,2)", "T.TEST(A1:A6,C1:C3,1,2)",
    "T.TEST(A1:A6,C1:C3,2,3)", "T.TEST(A1:A6,C1:C3,1,3)", "F.TEST(A1:A6,C1:C3)", "F.TEST(C1:C3,A1:A6)",
    "CHISQ.TEST(E1:F2,G1:H2)", "CHISQ.TEST(I1:I4,J1:J4)",
    "GAMMA(4.5)", "GAMMA(-1.5)", "GAMMA(0.5)", "GAMMA(10)", "GAMMA(0.1)", "GAMMA(-0.5)",
]
wb = openpyxl.Workbook()
ws = wb.active
for col, values in data.items():
    for i, v in enumerate(values, 1):
        ws[f"{col}{i}"] = v
for i, f in enumerate(formulas, 10):
    ws[f"A{i}"] = "=" + re.sub(r"\b(T\.TEST|F\.TEST|CHISQ\.TEST|GAMMA)\(", r"_xlfn.\1(", f)
src = OUT / "ref.xlsx"
wb.save(src)
for old in OUT.glob("ref*.csv"):
    old.unlink()
subprocess.run([so, "-env:UserInstallation=file:///E:/VNOffice-profiles/claude-lo-prof", "--headless", "--convert-to", "csv", "--outdir", str(OUT), str(src)], capture_output=True, timeout=240)
rows_csv = list(csv.reader((OUT / "ref.csv").open(encoding="utf-8")))
for f, row in zip(formulas, rows_csv[9:]):
    print(f"{f}\t{row[0]}")
